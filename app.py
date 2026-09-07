from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import streamlit as st
except ModuleNotFoundError:  # 核心邏輯可在未安裝 Streamlit 時測試
    st = None


APP_TITLE = "生管排程 × 製令缺料 × 客戶領料補料"
APP_VERSION = "V4－三表料號補料版"
WO_PATTERN = re.compile(r"(\d{2}[A-Z]{1,3}\d{4})-(\d{2})(?:\s*[~～]\s*(\d{2}))?", re.I)


ALIASES = {
    "MATERIAL": ["材料品號", "料號", "品號"],
    "NAME": ["品名", "材料名稱"],
    "SPEC": ["規格"],
    "RECEIPT_QTY": ["領料數量", "實際領料數量", "數量"],
    "RECEIPT_NO": ["領料單號", "單據號碼"],
    "UNIT": ["單位"],
    "NOTE": ["備註"],
    "SEQ": ["序號", "項次"],
    "SHORT_QTY": ["欠料數量", "缺料數量"],
    "STOCK": ["現有庫存", "庫存數量"],
    "URGENT": ["急料"],
    "MO_NUMBER": ["製令編號", "製令單號"],
    "WO": ["製令", "專案代號"],
    "ISSUE_DATE": ["發料日", "預計發料日"],
    "ENTRY_DATE": ["入庫日", "預計入庫日"],
    "CUSTOMER": ["客戶", "客戶簡稱"],
    "CATEGORY": ["CATEGORY", "類別"],
    "LOCATION": ["組立地點", "組立/地點"],
    "PROGRESS": ["組立進度", "進度"],
}


def norm_header(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return re.sub(r"[\s\n\r\t_（）()／/\-]+", "", str(value).strip().upper())


NORMALIZED_ALIASES = {
    key: {norm_header(alias) for alias in values} for key, values in ALIASES.items()
}


def clean_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def norm_material(value: Any) -> str:
    return re.sub(r"\s+", "", clean_text(value)).upper()


def extract_wos(value: Any) -> list[str]:
    text = clean_text(value).upper()
    found: list[str] = []
    for match in WO_PATTERN.finditer(text):
        base, start_text, end_text = match.groups()
        start = int(start_text)
        if end_text:
            end = int(end_text)
            if start <= end <= 99 and end - start <= 50:
                found.extend(f"{base}-{number:02d}" for number in range(start, end + 1))
                continue
        found.append(f"{base}-{start:02d}")
    return list(dict.fromkeys(found))


def to_number(value: Any) -> float:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return 0.0
    number = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    return 0.0 if pd.isna(number) else float(number)


def to_date(value: Any) -> pd.Timestamp | pd.NaT:
    if value is None or clean_text(value) in {"", "--", "-", "TBD", "NAN"}:
        return pd.NaT
    return pd.to_datetime(value, errors="coerce")


def find_column(df: pd.DataFrame, standard: str, required: bool = True) -> str | None:
    aliases = NORMALIZED_ALIASES[standard]
    for col in df.columns:
        if norm_header(col) in aliases:
            return col
    if required:
        readable = "／".join(ALIASES[standard])
        raise ValueError(f"找不到必要欄位：{readable}")
    return None


def detect_sheet_and_header(file_bytes: bytes, required: list[str], preferred: list[str] | None = None) -> tuple[str, int]:
    excel = pd.ExcelFile(io.BytesIO(file_bytes))
    sheets = excel.sheet_names
    ordered = []
    for name in (preferred or []):
        if name in sheets and name not in ordered:
            ordered.append(name)
    ordered.extend(name for name in sheets if name not in ordered)

    best: tuple[int, str, int] | None = None
    for sheet in ordered:
        preview = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, header=None, nrows=15, dtype=object)
        for row_idx in range(len(preview)):
            labels = {norm_header(v) for v in preview.iloc[row_idx].tolist()}
            score = sum(bool(labels & NORMALIZED_ALIASES[item]) for item in required)
            candidate = (score, sheet, row_idx)
            if best is None or candidate[0] > best[0]:
                best = candidate
            if score == len(required):
                return sheet, row_idx
    if not best or best[0] < len(required):
        missing = "、".join(required)
        raise ValueError(f"找不到符合欄位的工作表：{missing}")
    return best[1], best[2]


def read_auto(file_bytes: bytes, required: list[str], preferred: list[str] | None = None) -> tuple[pd.DataFrame, str]:
    sheet, header_row = detect_sheet_and_header(file_bytes, required, preferred)
    df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, header=header_row, dtype=object)
    df = df.dropna(how="all").copy()
    return df, sheet


def prepare_receipt(file_bytes: bytes) -> tuple[pd.DataFrame, str]:
    df, sheet = read_auto(file_bytes, ["MATERIAL", "RECEIPT_QTY", "NOTE"], ["單身資料"])
    receipt_no = ""
    try:
        header_df, _ = read_auto(file_bytes, ["RECEIPT_NO"], ["單頭資料"])
        receipt_no_col = find_column(header_df, "RECEIPT_NO")
        receipt_no_values = header_df[receipt_no_col].dropna().map(clean_text)
        if not receipt_no_values.empty:
            receipt_no = receipt_no_values.iloc[0]
    except (ValueError, KeyError):
        pass
    material_col = find_column(df, "MATERIAL")
    qty_col = find_column(df, "RECEIPT_QTY")
    note_col = find_column(df, "NOTE")
    name_col = find_column(df, "NAME", False)
    spec_col = find_column(df, "SPEC", False)
    unit_col = find_column(df, "UNIT", False)
    seq_col = find_column(df, "SEQ", False)

    result = pd.DataFrame({
        "領料單號": receipt_no,
        "領料序號": df[seq_col].map(clean_text) if seq_col else [str(i + 1) for i in range(len(df))],
        "材料品號": df[material_col].map(clean_text),
        "品名": df[name_col].map(clean_text) if name_col else "",
        "規格": df[spec_col].map(clean_text) if spec_col else "",
        "領料數量": df[qty_col].map(to_number),
        "單位": df[unit_col].map(clean_text) if unit_col else "",
        "領料備註": df[note_col].map(clean_text),
    })
    result["_材料"] = result["材料品號"].map(norm_material)
    result["_製令列表"] = result["領料備註"].map(extract_wos)
    result = result[(result["_材料"] != "") & (result["領料數量"] > 0)].reset_index(drop=True)
    return result, sheet


def prepare_shortage(file_bytes: bytes) -> tuple[pd.DataFrame, str]:
    df, sheet = read_auto(file_bytes, ["MATERIAL", "SHORT_QTY", "NOTE"])
    material_col = find_column(df, "MATERIAL")
    shortage_col = find_column(df, "SHORT_QTY")
    note_col = find_column(df, "NOTE")
    name_col = find_column(df, "NAME", False)
    spec_col = find_column(df, "SPEC", False)
    stock_col = find_column(df, "STOCK", False)
    urgent_col = find_column(df, "URGENT", False)
    mo_col = find_column(df, "MO_NUMBER", False)

    result = pd.DataFrame({
        "材料品號": df[material_col].map(clean_text),
        "品名": df[name_col].map(clean_text) if name_col else "",
        "規格": df[spec_col].map(clean_text) if spec_col else "",
        "製令編號": df[mo_col].map(clean_text) if mo_col else "",
        "急料": df[urgent_col].map(clean_text) if urgent_col else "",
        "欠料數量": df[shortage_col].map(to_number),
        "現有庫存": df[stock_col].map(to_number) if stock_col else 0.0,
        "欠料備註": df[note_col].map(clean_text),
    })
    result["_材料"] = result["材料品號"].map(norm_material)
    result["_製令列表"] = result["欠料備註"].map(extract_wos)
    result = result.explode("_製令列表", ignore_index=True).rename(columns={"_製令列表": "製令"})
    result["製令"] = result["製令"].fillna("").map(clean_text)
    result = result[(result["_材料"] != "") & (result["製令"] != "") & (result["欠料數量"] > 0)].reset_index(drop=True)
    result["_欠料列"] = result.index + 2
    return result, sheet


def prepare_schedule(file_bytes: bytes) -> tuple[pd.DataFrame, str]:
    df, sheet = read_auto(file_bytes, ["WO", "ISSUE_DATE"], ["2026排程"])
    wo_col = find_column(df, "WO")
    issue_col = find_column(df, "ISSUE_DATE")
    customer_col = find_column(df, "CUSTOMER", False)
    category_col = find_column(df, "CATEGORY", False)
    location_col = find_column(df, "LOCATION", False)
    entry_col = find_column(df, "ENTRY_DATE", False)
    progress_col = find_column(df, "PROGRESS", False)

    result = pd.DataFrame({
        "製令原值": df[wo_col].map(clean_text),
        "客戶": df[customer_col].map(clean_text) if customer_col else "",
        "Category": df[category_col].map(clean_text) if category_col else "",
        "組立地點": df[location_col].map(clean_text) if location_col else "",
        "組立進度": df[progress_col].map(clean_text) if progress_col else "",
        "發料日": df[issue_col].map(to_date),
        "入庫日": df[entry_col].map(to_date) if entry_col else pd.NaT,
    })
    result["_製令列表"] = result["製令原值"].map(extract_wos)
    result = result.explode("_製令列表", ignore_index=True).rename(columns={"_製令列表": "製令"})
    result["製令"] = result["製令"].fillna("").map(clean_text)
    result = result[result["製令"] != ""].copy()
    result["_有日期"] = result["發料日"].notna().astype(int)
    result = result.sort_values(["製令", "_有日期", "發料日"], ascending=[True, False, True], na_position="last")
    result = result.drop_duplicates("製令", keep="first").drop(columns=["_有日期"])
    return result.reset_index(drop=True), sheet


def priority_info(issue_date: Any, base_date: date) -> tuple[int, str, int | None]:
    if pd.isna(issue_date):
        return 4, "4. 排程未找到", None
    issued = pd.Timestamp(issue_date).date()
    delta = (base_date - issued).days
    if delta > 0:
        return 1, "1. 已過發料日", delta
    if delta == 0:
        return 2, "2. 今日發料", 0
    return 3, "3. 發料日未到", delta


def analyze(
    receipt: pd.DataFrame,
    shortage: pd.DataFrame,
    schedule: pd.DataFrame,
    base_date: date,
) -> dict[str, Any]:
    schedule_map = schedule.set_index("製令").to_dict("index")
    receipt_materials = set(receipt["_材料"])
    shortage_status = shortage.copy().reset_index(drop=True)

    for col in ["客戶", "Category", "組立地點", "組立進度", "發料日", "入庫日"]:
        shortage_status[col] = shortage_status["製令"].map(
            lambda wo: schedule_map.get(wo, {}).get(col, pd.NaT if "日" in col else "")
        )
    all_infos = shortage_status["發料日"].map(lambda value: priority_info(value, base_date))
    shortage_status["_優先群組"] = all_infos.map(lambda x: x[0])
    shortage_status["排程狀態"] = all_infos.map(lambda x: x[1])
    shortage_status["逾期天數"] = all_infos.map(lambda x: x[2] if x[2] is not None and x[2] >= 0 else 0)
    shortage_status["距發料日天數"] = all_infos.map(lambda x: abs(x[2]) if x[2] is not None and x[2] < 0 else 0)
    shortage_status["_日期排序"] = shortage_status["發料日"].fillna(pd.Timestamp.max)
    shortage_status = shortage_status.sort_values(
        ["_優先群組", "_日期排序", "製令", "_材料", "_欠料列"], kind="stable"
    ).reset_index(drop=True)

    detail = shortage_status[
        shortage_status["_材料"].isin(receipt_materials) & (shortage_status["_優先群組"] == 1)
    ].copy()
    detail["優先狀態"] = detail["排程狀態"]
    detail["_日期排序"] = detail["發料日"].fillna(pd.Timestamp.max)
    detail = detail.sort_values(["_日期排序", "製令", "_材料", "_欠料列"], kind="stable").reset_index(drop=True)

    ordered_wos = list(dict.fromkeys(detail["製令"].tolist()))
    wo_rank = {wo: rank for rank, wo in enumerate(ordered_wos, start=1)}
    detail["製令優先序"] = detail["製令"].map(wo_rank)
    detail["建議補料數量"] = 0.0
    detail["領料來源序號"] = ""

    by_material: dict[str, list[int]] = defaultdict(list)
    for idx, row in detail.iterrows():
        by_material[row["_材料"]].append(idx)

    allocation_log: list[dict[str, Any]] = []
    unmatched_receipt: list[dict[str, Any]] = []

    for receipt_idx, receipt_row in receipt.iterrows():
        material = receipt_row["_材料"]
        remaining = float(receipt_row["領料數量"])
        candidate_indices = list(by_material.get(material, []))
        candidate_indices = sorted(candidate_indices, key=lambda idx: (
            detail.at[idx, "_日期排序"], detail.at[idx, "製令"], detail.at[idx, "_欠料列"]
        ))

        for detail_idx in candidate_indices:
            if remaining <= 0:
                break
            needed = max(float(detail.at[detail_idx, "欠料數量"]) - float(detail.at[detail_idx, "建議補料數量"]), 0)
            allocated = min(remaining, needed)
            if allocated <= 0:
                continue
            detail.at[detail_idx, "建議補料數量"] += allocated
            previous = clean_text(detail.at[detail_idx, "領料來源序號"])
            seq = clean_text(receipt_row["領料序號"])
            detail.at[detail_idx, "領料來源序號"] = f"{previous}、{seq}" if previous else seq
            remaining -= allocated
            allocation_log.append({
                "補料次序": 0,
                "領料單號": receipt_row["領料單號"],
                "製令優先序": int(detail.at[detail_idx, "製令優先序"]),
                "優先狀態": detail.at[detail_idx, "優先狀態"],
                "製令": detail.at[detail_idx, "製令"],
                "發料日": detail.at[detail_idx, "發料日"],
                "材料品號": receipt_row["材料品號"],
                "品名": receipt_row["品名"],
                "領料序號": receipt_row["領料序號"],
                "領料數量": receipt_row["領料數量"],
                "該列欠料數量": detail.at[detail_idx, "欠料數量"],
                "建議補料數量": allocated,
                "領料備註": receipt_row["領料備註"],
            })

        if remaining > 0:
            reason = "已過發料日的製令缺料中查無此材料品號" if not candidate_indices else "領料數量超過已過發料日製令的欠料數量"
            unmatched_receipt.append({
                "領料單號": receipt_row["領料單號"], "領料序號": receipt_row["領料序號"], "材料品號": receipt_row["材料品號"],
                "品名": receipt_row["品名"], "領料數量": receipt_row["領料數量"], "未分配數量": remaining,
                "領料備註": receipt_row["領料備註"], "原因": reason,
            })

    detail["補料判定"] = detail.apply(
        lambda row: "可由本次領料補料" if row["建議補料數量"] > 0
        else "領料量已優先分配給更早逾期製令",
        axis=1,
    )

    allocation_df = pd.DataFrame(allocation_log)
    if not allocation_df.empty:
        allocation_df = allocation_df.sort_values(["製令優先序", "發料日", "製令", "材料品號"], kind="stable").reset_index(drop=True)
        allocation_df["補料次序"] = range(1, len(allocation_df) + 1)

    summary = shortage_status.groupby("製令", as_index=False).agg(
        缺料品項數=("材料品號", "size"),
        欠料總數量=("欠料數量", "sum"),
    )
    matched_summary = detail.groupby("製令", as_index=False).agg(
        可比對領料品項數=("材料品號", "size"),
        可補料品項數=("建議補料數量", lambda s: int((s > 0).sum())),
        建議補料總數量=("建議補料數量", "sum"),
    )
    summary = summary.merge(matched_summary, on="製令", how="left")
    for col in ["可比對領料品項數", "可補料品項數", "建議補料總數量"]:
        summary[col] = summary[col].fillna(0)
    for col in ["客戶", "Category", "組立地點", "發料日", "入庫日"]:
        summary[col] = summary["製令"].map(lambda wo: schedule_map.get(wo, {}).get(col, pd.NaT if "日" in col else ""))
    summary["排程狀態"] = summary["發料日"].map(lambda value: priority_info(value, base_date)[1])
    summary["_優先群組"] = summary["發料日"].map(lambda value: priority_info(value, base_date)[0])
    summary["_日期排序"] = summary["發料日"].fillna(pd.Timestamp.max)
    summary = summary.sort_values(["_優先群組", "_日期排序", "製令"], kind="stable").reset_index(drop=True)
    summary.insert(0, "製令排序", range(1, len(summary) + 1))
    summary = summary.drop(columns=["_優先群組", "_日期排序"])

    visible_columns = [
        "製令優先序", "優先狀態", "逾期天數", "製令", "客戶", "Category", "組立地點",
        "發料日", "入庫日", "材料品號", "品名", "規格", "製令編號", "急料", "欠料數量", "現有庫存",
        "建議補料數量", "補料判定", "領料來源序號", "欠料備註", "_欠料列",
    ]
    detail_visible = detail[visible_columns].rename(columns={"_欠料列": "欠料表列號"})
    shortage_status_columns = [
        "排程狀態", "逾期天數", "距發料日天數", "製令", "客戶", "Category", "組立地點",
        "發料日", "入庫日", "材料品號", "品名", "規格", "製令編號", "急料", "欠料數量",
        "現有庫存", "欠料備註", "_欠料列",
    ]
    shortage_status_visible = shortage_status[shortage_status_columns].rename(columns={"_欠料列": "欠料表列號"})
    unmatched_df = pd.DataFrame(unmatched_receipt)
    stats = {
        "領料筆數": len(receipt),
        "領料單號": "、".join(receipt["領料單號"].dropna().astype(str).unique().tolist()),
        "缺料製令數": shortage_status["製令"].nunique(),
        "逾期製令數": len(ordered_wos),
        "逾期缺料筆數": len(detail_visible),
        "可補料明細數": len(allocation_df),
        "未分配領料筆數": len(unmatched_df),
    }
    return {
        "allocation": allocation_df,
        "detail": detail_visible,
        "shortage_status": shortage_status_visible,
        "summary": summary,
        "unmatched": unmatched_df,
        "stats": stats,
    }


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PRIORITY_FILLS = {
    "1. 已過發料日": PatternFill("solid", fgColor="F4CCCC"),
    "2. 今日發料": PatternFill("solid", fgColor="FFF2CC"),
    "3. 發料日未到": PatternFill("solid", fgColor="D9EAD3"),
    "4. 排程未找到": PatternFill("solid", fgColor="E7E6E6"),
}


def excel_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and pd.isna(value)) or value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item") and not isinstance(value, (str, bytes, date, datetime)):
        try:
            return value.item()
        except (AttributeError, ValueError):
            pass
    return value


def add_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title)
    data = df.copy()
    if data.empty and len(data.columns) == 0:
        data = pd.DataFrame({"結果": ["無資料"]})
    for col_idx, col in enumerate(data.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(col))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    priority_col = list(data.columns).index("優先狀態") + 1 if "優先狀態" in data.columns else None
    date_cols = {i + 1 for i, col in enumerate(data.columns) if "日期" in str(col) or str(col) in {"發料日", "入庫日"}}
    for row_idx, values in enumerate(data.itertuples(index=False, name=None), start=2):
        fill = PRIORITY_FILLS.get(clean_text(values[priority_col - 1])) if priority_col else None
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=excel_value(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx in date_cols and isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd"
            if fill:
                cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for col_idx, col in enumerate(data.columns, start=1):
        samples = [str(col)] + [clean_text(v) for v in data.iloc[:100, col_idx - 1].tolist()]
        width = min(max(max((len(v) for v in samples), default=8) + 2, 10), 34)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def build_output_excel(result: dict[str, Any], base_date: date, sources: dict[str, str]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    info = pd.DataFrame({
        "項目": ["判斷基準日", "比對規則", "排序規則", "領料單工作表", "製令欠料工作表", "製程排程工作表"],
        "內容": [
            base_date.strftime("%Y/%m/%d"),
            "依領料單的材料品號，比對發料日早於判斷基準日的製令缺料；不限制領料備註中的製令。",
            "發料日逾期較久的製令優先；相同製令依材料品號排序，補料數量不超過欠料數量。",
            sources.get("receipt", ""), sources.get("shortage", ""), sources.get("schedule", ""),
        ],
    })
    add_sheet(wb, "使用說明", info)
    add_sheet(wb, "建議補料順序", result["allocation"])
    add_sheet(wb, "製令缺料彙總", result["summary"])
    add_sheet(wb, "逾期製令缺料比對", result["detail"])
    add_sheet(wb, "未分配領料", result["unmatched"])
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def main() -> None:
    if st is None:
        raise RuntimeError("尚未安裝 Streamlit，請執行 pip install -r requirements.txt。")

    st.set_page_config(page_title=APP_TITLE, page_icon="📦", layout="wide")
    st.title("📦 生管排程 × 製令缺料 × 客戶領料補料")
    st.success(f"目前版本：{APP_VERSION}")
    st.caption("生管排程確認發料日，製令缺料表確認缺料；客戶領回料件暫時不看製令，只依料號補到已過發料日的製令。")

    c1, c2, c3 = st.columns(3)
    with c1:
        schedule_file = st.file_uploader("① 上傳生管排程表（XLSX）", type=["xlsx"], key="schedule")
    with c2:
        shortage_file = st.file_uploader("② 上傳製令缺料表（XLSX）", type=["xlsx"], key="shortage")
    with c3:
        receipt_file = st.file_uploader("③ 上傳客戶端領料單（XLSX）", type=["xlsx"], key="receipt")

    base_date = st.date_input("判斷基準日", value=date.today())
    st.info("僅比對發料日早於判斷基準日的製令；今日及未來發料的製令不納入補料。")

    if not receipt_file or not shortage_file or not schedule_file:
        st.info("請依序上傳生管排程表、製令缺料表、客戶端領料單三份檔案。")
        return

    if st.button("開始比對並產生補料順序", type="primary", use_container_width=True):
        try:
            with st.spinner("正在讀取三份資料並計算補料優先順序，欠料表較大時請稍候…"):
                receipt, receipt_sheet = prepare_receipt(receipt_file.getvalue())
                shortage, shortage_sheet = prepare_shortage(shortage_file.getvalue())
                schedule, schedule_sheet = prepare_schedule(schedule_file.getvalue())
                result = analyze(receipt, shortage, schedule, base_date)
                output = build_output_excel(result, base_date, {
                    "receipt": receipt_sheet, "shortage": shortage_sheet, "schedule": schedule_sheet,
                })
                st.session_state["analysis_result"] = result
                st.session_state["analysis_output"] = output
                st.session_state["analysis_date"] = base_date
        except Exception as exc:
            st.error(f"處理失敗：{exc}")

    if "analysis_result" not in st.session_state:
        return

    result = st.session_state["analysis_result"]
    stats = result["stats"]
    m1, m2, m3, m4, m5, m6 = st.columns(6)
    m1.metric("領料單號", stats["領料單號"] or "未取得")
    m2.metric("缺料表製令", f"{stats['缺料製令數']} 張")
    m3.metric("已過發料日可比對製令", f"{stats['逾期製令數']} 張")
    m4.metric("料號相符的逾期缺料", f"{stats['逾期缺料筆數']} 筆")
    m5.metric("建議補料", f"{stats['可補料明細數']} 筆")
    m6.metric("未分配領料", f"{stats['未分配領料筆數']} 筆")

    tab1, tab2, tab3, tab4, tab5 = st.tabs([
        "建議補料順序", "逾期製令缺料比對", "全部製令缺料狀況", "製令缺料彙總", "未分配領料"
    ])
    with tab1:
        if result["allocation"].empty:
            st.warning("領料單的材料品號在已超過發料日的製令缺料中沒有可補料項目。")
        else:
            st.dataframe(result["allocation"], use_container_width=True, hide_index=True, height=520)
    with tab2:
        detail = result["detail"]
        selected_wo = st.multiselect("篩選逾期製令", sorted(detail["製令"].dropna().unique().tolist()))
        show = detail[detail["製令"].isin(selected_wo)] if selected_wo else detail
        st.dataframe(show.head(10000), use_container_width=True, hide_index=True, height=580)
        if len(show) > 10000:
            st.caption("畫面僅預覽前 10,000 筆，下載 Excel 仍包含全部資料。")
    with tab3:
        all_shortage = result["shortage_status"]
        all_selected_wo = st.multiselect(
            "選擇製令查看所有缺料料件",
            sorted(all_shortage["製令"].dropna().unique().tolist()),
            key="all_shortage_wo",
        )
        all_show = all_shortage[all_shortage["製令"].isin(all_selected_wo)] if all_selected_wo else all_shortage
        st.dataframe(all_show.head(10000), use_container_width=True, hide_index=True, height=580)
        if len(all_show) > 10000:
            st.caption("畫面僅預覽前 10,000 筆；可先選擇製令縮小範圍。")
    with tab4:
        st.dataframe(result["summary"], use_container_width=True, hide_index=True)
    with tab5:
        if result["unmatched"].empty:
            st.success("所有領料數量均已分配到指定製令缺料。")
        else:
            st.dataframe(result["unmatched"], use_container_width=True, hide_index=True, height=450)

    output_date = st.session_state["analysis_date"].strftime("%Y%m%d")
    st.download_button(
        "下載領料單逾期製令補料 Excel",
        data=st.session_state["analysis_output"],
        file_name=f"領料單逾期製令補料_{output_date}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )


if __name__ == "__main__":
    main()
